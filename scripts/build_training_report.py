from __future__ import annotations

import os
import re
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"E:\桌面\smart_finance_agent")
TEMPLATE = Path(r"E:\桌面\软件系统分析及设计综合实训--实训报告参考模板（附写作要求）.docx")
OUT_DIR = ROOT / "outputs"
OUT_DOCX = OUT_DIR / "software-system-analysis-design-report-final.docx"
FALLBACK_DOCX = OUT_DIR / "software-system-analysis-design-report-final-fixed.docx"

SCREENSHOTS = [
    ("登录页", Path(r"C:\TEMP\codex-clipboard-73fc509e-ba0b-469e-9851-ac139b2a945d.png"),
     "用户通过登录页进入系统，系统以暗色工作台风格呈现智财 Agent 的定位。"),
    ("新对话页", Path(r"C:\TEMP\codex-clipboard-ac9c23b0-9a65-4afe-ae84-546052b901d2.png"),
     "登录成功后进入新对话页面，用户可以直接向 Agent 发起记账、预算和消费分析请求。"),
    ("财务画像页", Path(r"C:\TEMP\codex-clipboard-e7490f25-98be-42c1-9091-61e71cff8d0e.png"),
     "财务画像用于维护收入、固定支出、储蓄目标、风险偏好和长期指令。"),
    ("Agent 审计页", Path(r"C:\TEMP\codex-clipboard-1e8a08df-1c58-44f5-9771-6b6211db3eb3.png"),
     "Agent 审计页统一展示反思、待确认动作、周期任务和运行记录。"),
    ("周期任务页", Path(r"C:\TEMP\codex-clipboard-584b5923-8a86-45cb-94e2-1285c227488b.png"),
     "周期任务页用于管理自动复盘、预算检查和监控类任务。"),
    ("Agent 技能页", Path(r"C:\TEMP\codex-clipboard-5c9669d4-9935-4ecd-8d5d-d8b9b69c900e.png"),
     "Agent 技能页展示可启停的内置 Skill，包括搜索、分类查询、周期任务创建和自定义 Skill 草稿。"),
    ("统计页", Path(r"C:\TEMP\codex-clipboard-275b408c-594e-41d1-9a25-fdd6a06f8aea.png"),
     "统计页按日、月、年展示收入、支出和结余趋势。"),
    ("账单导入页", Path(r"C:\TEMP\codex-clipboard-19e4505e-6e29-4506-bd80-3387707a9852.png"),
     "账单导入页支持上传微信、支付宝或银行卡流水截图，由多模态模型识别候选交易。"),
    ("消费记录页", Path(r"C:\TEMP\codex-clipboard-5e75d57b-4e78-40a1-b4fc-6be0331d4996.png"),
     "消费记录页支持分类、日期筛选以及记录编辑和删除。"),
    ("个性化设置页", Path(r"C:\TEMP\codex-clipboard-5e94e077-6e8b-4279-b164-88fde7e0a9b7.png"),
     "个性化设置页支持切换主题、主题色、显示密度和内容宽度。"),
]

UML_IMAGES = {
    "图2-1 业务用例图": Path(r"C:\TEMP\codex-clipboard-5f735bae-bf69-46c2-94fa-f11033b107d8.png"),
    "图2-2 账单导入业务活动图": Path(r"C:\TEMP\codex-clipboard-2609fb98-8c6b-48ed-bd38-b57cbf540008.png"),
    "图3-1 系统用例图": Path(r"C:\TEMP\codex-clipboard-7ec74fe0-4190-45fe-94e9-78ef21ddda9f.png"),
    "图4-1 备选构架图": Path(r"C:\TEMP\codex-clipboard-e2160a0e-e7b0-4db6-9267-ce484912f164.png"),
    "图4-2 智能财务问答活动图": Path(r"C:\TEMP\codex-clipboard-53dfd79e-a63d-4d5e-a81f-aa6af8a8e3f6.png"),
    "图4-3 智能财务问答 VOPC 类图": Path(r"C:\TEMP\codex-clipboard-9aeba0a9-e1bf-45d9-a1ff-34937f9b2b69.png"),
    "图4-4 账单导入确认活动图": Path(r"C:\TEMP\codex-clipboard-2e3c66b4-5bdf-42f4-9b9b-4d760f0e2ed0.png"),
    "图4-5 账单导入 VOPC 类图": Path(r"C:\TEMP\codex-clipboard-ff1f089c-9a63-490e-afdc-1030895e6118.png"),
    "图4-6 实体类类图": Path(r"C:\TEMP\codex-clipboard-3e5e15aa-f18b-47e4-9f6c-1c9caf5b4875.png"),
    "图5-1 功能模块图": Path(r"C:\TEMP\codex-clipboard-aa46740c-0f64-490b-b015-695c5da2094b.png"),
    "图5-2 系统架构分层图": Path(r"C:\TEMP\codex-clipboard-60271e8d-c713-4aa3-bb55-77d92991df9a.png"),
    "图5-3 详细设计类图": Path(r"C:\TEMP\codex-clipboard-d694dc36-5e40-43f8-926e-879864f9016d.png"),
    "图5-4 数据库 ER 图": Path(r"C:\TEMP\codex-clipboard-256292ef-0bfa-4c36-ae53-99bb2dd55798.png"),
}


def set_run_font(run, font="宋体", size=None, bold=None, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, fill=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(str(text))
    set_run_font(run, font="宋体", size=10.5, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if fill:
        set_cell_shading(cell, fill)


def set_table_widths(table, widths_cm):
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            row.cells[idx].width = Cm(width)


def clear_document(doc):
    body = doc._body._element
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def block_text(block):
    return "".join(t.text or "" for t in block.iter(qn("w:t"))).strip()


def trim_template_body_from(doc, start_text):
    body = doc._body._element
    children = list(body)
    start_idx = None
    for idx, child in enumerate(children):
        if block_text(child) == start_text:
            start_idx = idx
            break
    if start_idx is None:
        clear_document(doc)
        return False
    for child in children[start_idx:]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return True


def fill_template_cover(doc):
    if not doc.tables:
        return
    table = doc.tables[0]
    values = {
        "项目名称": "智财 Agent",
        "专    业": "软件工程",
        "班    级": "",
        "姓    名": "",
        "完成时间": "2026年7月",
    }
    for row in table.rows:
        if len(row.cells) < 2:
            continue
        key = row.cells[0].text.strip()
        if key in values:
            row.cells[1].text = values[key]
            for p in row.cells[1].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_run_font(run, size=10.5)


def setup_document(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(4)

    for name, size in [("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)]:
        style = styles[name]
        style.font.name = "黑体"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)


def add_p(doc, text="", bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, size=10.5, bold=True)
        rest = text[len(bold_prefix):]
        if rest:
            run = p.add_run(rest)
            set_run_font(run, size=10.5)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5)
    return p


def add_heading(doc, text, level):
    # The report template already applies automatic numbering to heading styles.
    # Strip manual chapter numbers to avoid visual duplicates such as "1  1 系统概述".
    text = re.sub(r"^\s*\d+(?:\.\d+)*\s+", "", text)
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_run_font(run, font="黑体", size={1: 16, 2: 14, 3: 12}.get(level, 11), bold=True)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, size=9, color="555555")
    return p


def add_table(doc, headers, rows, widths_cm=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, header in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], header, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, fill="E8EEF5")
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            align = WD_ALIGN_PARAGRAPH.CENTER if len(str(val)) <= 12 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cells[i], val, align=align)
    if widths_cm:
        set_table_widths(table, widths_cm)
    doc.add_paragraph()
    return table


def add_placeholder(doc, caption, description):
    image_path = UML_IMAGES.get(caption)
    if image_path and image_path.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(image_path), width=Inches(6.25))
        add_caption(doc, caption)
        return

    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F2F4F7")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(18)
    run = p.add_run(f"{caption}\n（待插入 PlantUML 渲染图片）\n{description}")
    set_run_font(run, font="宋体", size=10, bold=True, color="555555")
    add_caption(doc, caption)


def add_code_block(doc, title, code):
    p = doc.add_paragraph()
    run = p.add_run(title)
    set_run_font(run, font="黑体", size=10.5, bold=True)
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F7F7F7")
    cp = cell.paragraphs[0]
    for line in code.strip().splitlines():
        r = cp.add_run(line.rstrip() + "\n")
        set_run_font(r, font="Consolas", size=8.5)
    doc.add_paragraph()


def add_cover(doc):
    for text, size, bold in [
        ("《软件系统分析及设计综合实训》", 20, True),
        ("实训报告", 24, True),
        ("2025-2026学年 第2学期", 14, False),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(text)
        set_run_font(run, font="黑体" if bold else "宋体", size=size, bold=bold)

    doc.add_paragraph()
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    items = [
        ("项目名称", "智财 Agent"),
        ("专    业", "软件工程"),
        ("班    级", ""),
        ("姓    名", ""),
        ("完成时间", "2026年7月"),
    ]
    for i, (k, v) in enumerate(items):
        set_cell_text(table.cell(i, 0), k, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_text(table.cell(i, 1), v, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_table_widths(table, [4.2, 8.6])
    doc.add_page_break()


def add_toc(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("目  录")
    set_run_font(run, font="黑体", size=16, bold=True)
    entries = [
        "1 系统概述",
        "1.1 研究背景及意义",
        "1.2 核心业务",
        "1.3 系统拟解决的主要问题",
        "2 业务建模",
        "2.1 建立业务用例模型",
        "2.2 分析业务用例",
        "3 用例建模",
        "3.1 识别参与者",
        "3.2 识别用例",
        "3.3 用例图和用例文档",
        "4 用例分析",
        "4.1 备选构架",
        "4.2 活动图和 VOPC 类图",
        "4.3 实体类类图",
        "5 系统设计",
        "5.1 功能模块",
        "5.2 技术选型与运行环境",
        "5.3 系统架构",
        "5.4 接口设计",
        "5.5 类设计",
        "5.6 数据库设计",
        "6 编码实现",
        "6.1 项目结构",
        "6.2 关键代码",
        "6.3 实现效果",
        "7 总结与体会",
    ]
    for entry in entries:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6 if entry.count(".") else 0)
        run = p.add_run(entry)
        set_run_font(run, size=10.5)
    add_p(doc, "说明：最终提交前可在 Word 中使用“引用-目录-更新目录”生成自动页码。")
    doc.add_page_break()


def add_system_overview(doc):
    add_heading(doc, "1 系统概述", 1)
    add_heading(doc, "1.1 研究背景及意义", 2)
    add_p(doc, "随着移动支付、线上消费和个人理财场景的普及，个人财务数据越来越分散在微信、支付宝、银行卡流水、手工记账和临时备忘中。普通用户往往能够记录某一笔消费，却难以持续形成完整的收支结构、预算执行情况和长期财务画像。传统记账软件多停留在“录入-统计”的工具层面，用户需要手动分类、手动复盘、手动理解图表，使用门槛和持续维护成本较高。")
    add_p(doc, "本项目面向个人用户设计并实现“智财 Agent”个人智能财务代理系统，将账单导入、消费记录、统计分析、财务画像和智能对话结合起来，在基础财务管理能力之上引入 Agent 化工作流。系统不仅提供传统的收支记录和预算统计，还支持基于大模型的财务问答、RAG 知识增强、ReAct 工具调用、长期记忆、反思建议、周期任务和待确认动作，从而把分散的个人财务操作组织成一个可追踪、可审计、可扩展的智能工作台。")
    add_p(doc, "系统建设的意义主要体现在三个方面：第一，降低个人财务管理门槛，使用户可以通过对话完成记录、查询和分析；第二，提升财务数据的连续性和可解释性，通过图表、画像和 Agent 运行轨迹帮助用户理解自身消费行为；第三，探索大模型在垂直财务场景中的安全落地方式，对关键写操作采用确认机制，避免 AI 直接修改用户数据。")

    add_heading(doc, "1.2 核心业务", 2)
    add_table(doc, ["核心业务", "发起角色", "处理数据", "产生结果"], [
        ("用户认证与工作台进入", "普通用户", "用户名、密码、JWT 令牌", "完成登录注册并进入个人财务工作台"),
        ("收支记录管理", "普通用户", "金额、类型、分类、日期、备注", "形成可查询、可统计的交易记录"),
        ("财务画像维护", "普通用户", "收入、固定支出、储蓄目标、风险偏好、长期指令", "为预算建议和 Agent 回答提供长期背景"),
        ("账单截图导入", "普通用户", "微信、支付宝或银行卡流水截图", "识别候选交易，用户确认后写入正式记录"),
        ("智能财务对话", "普通用户", "自然语言问题、历史上下文、RAG 知识和工具结果", "输出财务分析、预算建议或待确认动作"),
        ("Agent 运行与审计", "普通用户/管理员", "traceId、工具调用、反思、周期任务、待确认动作", "实现可追踪、可审计的智能执行链路"),
    ], [3.2, 2.4, 5.2, 4.2])

    add_heading(doc, "1.3 系统拟解决的主要问题", 2)
    problems = [
        "个人财务数据分散，用户难以形成统一视图。系统通过交易记录、分类、预算、统计和财务画像整合个人财务数据。",
        "手工记账成本高，账单截图难以直接利用。系统通过多模态账单识别生成候选交易，并要求用户确认后入库。",
        "传统统计图表缺少解释能力。系统提供 AI 对话、RAG 知识和工具调用，使用户可以用自然语言理解消费趋势和预算风险。",
        "AI 写操作存在误操作风险。系统通过待确认动作机制，将记账、预算设置、自定义 Skill 安装等关键操作交由用户确认。",
        "大模型回答缺少可追踪性。系统记录 Agent run、Agent run step、analysis record 和 skill invocation，使执行过程可审计。",
    ]
    for i, item in enumerate(problems, 1):
        add_p(doc, f"（{i}）{item}")
    add_p(doc, "系统战略目标是形成一个面向个人用户的智能财务工作台，使财务数据可录入、可统计、可解释、可复盘；质量目标包括功能完整性、数据安全性、操作可确认性、模块可扩展性和运行可追踪性。")


def add_business_model(doc):
    add_heading(doc, "2 业务建模", 1)
    add_heading(doc, "2.1 建立业务用例模型", 2)
    add_p(doc, "本系统的现实业务围绕个人财务管理展开，用户首先完成登录并维护个人财务画像，然后通过手工新增或账单导入形成交易数据。系统根据交易记录生成统计结果，并在用户发起咨询或周期任务触发时，结合画像、历史记录、财务知识和可用工具给出建议。")
    add_placeholder(doc, "图2-1 业务用例图", "展示用户完成账单导入、记账、统计分析、财务咨询和周期复盘的业务目标。")
    add_heading(doc, "2.2 分析业务用例", 2)
    add_p(doc, "以“账单导入并确认入库”为例，业务开始于用户上传账单截图。系统调用多模态识别能力判断账单来源并抽取候选交易，用户核对金额、日期、分类和备注后确认入库，系统最终生成正式交易记录并进入统计分析范围。")
    add_placeholder(doc, "图2-2 账单导入业务活动图", "展示上传截图、AI 识别、生成候选交易、用户确认和写入交易记录的业务流程。")


def add_use_cases(doc):
    add_heading(doc, "3 用例建模", 1)
    add_heading(doc, "3.1 识别参与者", 2)
    add_table(doc, ["参与者", "职责说明"], [
        ("普通用户", "使用系统完成登录、维护财务画像、管理收支记录、导入账单、查看统计结果、与 Agent 对话并确认关键动作。"),
        ("系统管理员", "从管理和演示角度查看系统运行状态、Agent 审计信息和异常任务，辅助维护系统。"),
        ("AI 大模型服务", "为智能问答、ReAct 推理、反思建议和多模态账单识别提供模型能力。"),
        ("外部搜索服务", "在需要实时财经、汇率、市场新闻等外部信息时提供检索结果。"),
        ("数据库系统", "持久化用户、交易、预算、画像、会话、Agent 运行轨迹和待确认动作等数据。"),
    ], [3, 12])

    add_heading(doc, "3.2 识别用例", 2)
    add_table(doc, ["用例编号", "用例名称", "业务目标"], [
        ("UC01", "注册登录", "识别用户身份，建立个人数据边界。"),
        ("UC02", "维护财务画像", "维护收入、支出、储蓄目标、风险偏好和长期指令。"),
        ("UC03", "管理消费记录", "新增、查询、筛选、编辑和删除个人收支记录。"),
        ("UC04", "导入账单截图", "从账单截图中识别候选交易并经用户确认入库。"),
        ("UC05", "查看统计分析", "按日、月、年查看收入、支出、结余和趋势图表。"),
        ("UC06", "智能财务问答", "通过自然语言获得记账、预算、消费分析和理财建议。"),
        ("UC07", "管理 Agent 技能", "查看、启停和安装 Agent 可使用的 Skill。"),
        ("UC08", "管理周期任务", "创建、启停、重试和查看自动复盘或预算检查任务。"),
        ("UC09", "确认待执行动作", "对 AI 生成的写操作进行确认或取消。"),
        ("UC10", "查看 Agent 审计", "查看反思、运行轨迹、周期任务和异常事件。"),
    ], [2.2, 3.2, 9])

    add_heading(doc, "3.3 用例图和用例文档", 2)
    add_heading(doc, "3.3.1 系统用例图", 3)
    add_placeholder(doc, "图3-1 系统用例图", "展示普通用户、管理员、AI 大模型服务、外部搜索服务与系统核心用例之间的关系。")
    add_heading(doc, "3.3.2 用例文档", 3)
    use_case_docs = [
        ("UC03", "管理消费记录", "用户维护个人收入和支出记录。", "普通用户", "已登录，存在有效 JWT。", "交易记录被新增、修改、删除或查询结果返回。", "用户进入消费记录页；系统展示筛选条件和记录列表；用户新增或编辑金额、类型、分类、日期、备注；系统校验并保存；用户可按类型、分类、日期过滤。", "金额为空或格式错误时提示校验失败；删除记录前需要用户确认。", "交易数据必须归属当前登录用户。"),
        ("UC04", "导入账单截图", "用户上传流水截图并确认候选交易。", "普通用户、AI 大模型服务", "用户已登录并选择有效图片文件。", "确认后的候选交易写入正式 transaction 表。", "用户上传截图；系统保存原始文件；调用多模态模型识别账单来源和交易信息；生成候选交易；用户核对并确认；系统写入交易记录。", "识别置信度低时标记为低置信；非账单图片不进入确认流程；用户可忽略候选交易。", "涉及姓名、卡号、订单号等隐私信息时建议用户先打码。"),
        ("UC06", "智能财务问答", "用户通过对话获得财务分析或执行建议。", "普通用户、AI 大模型服务、外部搜索服务", "用户已登录并输入问题。", "系统返回回答、traceId 和可审计运行步骤。", "用户输入问题；系统组装画像、历史、RAG 知识和可用工具；ReActAgentService 决策是否调用工具；ToolRegistry 执行工具；系统汇总结果并流式返回。", "模型调用失败时返回错误说明；写操作转为待确认动作而非直接执行。", "最多执行有限步数，避免无限循环；敏感数据不自动沉淀为长期记忆。"),
        ("UC08", "管理周期任务", "用户配置 Agent 自动复盘或预算检查任务。", "普通用户", "用户已登录并填写任务名称、执行内容和 cron 表达式。", "周期任务被保存，可启停并记录运行历史。", "用户进入周期任务页；新增任务；系统保存任务配置；调度器按时间触发；任务运行结果写入 Agent schedule run。", "cron 表达式错误时拒绝保存；任务失败时记录错误并可重试。", "周期任务执行仍遵守 Agent 安全边界。"),
    ]
    for uc in use_case_docs:
        rows = [
            ("用例编号", uc[0], "用例名", uc[1]),
            ("简要描述", uc[2], "参与者", uc[3]),
            ("前置条件", uc[4], "后置条件", uc[5]),
            ("基本事件流", uc[6], "", ""),
            ("备选事件流", uc[7], "", ""),
            ("补充约束", uc[8], "", ""),
        ]
        table = doc.add_table(rows=0, cols=4)
        table.style = "Table Grid"
        for row in rows:
            cells = table.add_row().cells
            for i, val in enumerate(row):
                set_cell_text(cells[i], val, bold=i in [0, 2], fill="F2F4F7" if i in [0, 2] else None)
        set_table_widths(table, [2.5, 5, 2.5, 5])
        doc.add_paragraph()


def add_analysis(doc):
    add_heading(doc, "4 用例分析", 1)
    add_heading(doc, "4.1 备选构架", 2)
    add_p(doc, "分析阶段将系统对象划分为边界类、控制类和实体类。边界类负责页面或接口交互，控制类负责协调业务流程，实体类对应需要持久化的业务数据。系统整体采用前后端分离和后端分层结构，前端视图调用 API，后端 Controller 接收请求，Service 处理业务逻辑，Mapper 完成数据库访问，Agent 模块负责上下文组装、ReAct 推理和工具调用。")
    add_table(doc, ["类型", "代表类或模块", "职责"], [
        ("边界类", "Login.vue、ChatView.vue、Transaction.vue、BillImport.vue、AgentSchedules.vue、各 Controller", "接收用户操作或 HTTP 请求，展示结果并传递参数。"),
        ("控制类", "ChatServiceImpl、ReActAgentService、BillImportServiceImpl、PendingActionServiceImpl、AgentScheduleServiceImpl", "组织业务流程、调用外部模型或工具、控制写操作确认。"),
        ("实体类", "User、Transaction、Budget、FinancialProfile、ChatConversation、AgentRun、AgentMemory、AgentSkill、PendingAction", "表达业务数据和持久化结构。"),
    ], [2.5, 6, 7])
    add_placeholder(doc, "图4-1 备选构架图", "展示边界类、控制类、实体类在表现层、业务层、持久层和 Agent 能力层中的分布。")

    add_heading(doc, "4.2 活动图和 VOPC 类图", 2)
    add_p(doc, "本节选择“智能财务问答”和“账单导入确认”两个核心用例进行分析。智能财务问答体现 Agent 的上下文组装、工具调用和审计轨迹；账单导入确认体现多模态识别结果必须经用户确认后才写入正式交易记录。")
    add_placeholder(doc, "图4-2 智能财务问答活动图", "展示用户提问、上下文组装、RAG 检索、ReAct 工具调用、流式返回和运行轨迹持久化过程。")
    add_placeholder(doc, "图4-3 智能财务问答 VOPC 类图", "展示 ChatView、ChatController、ChatServiceImpl、ReActAgentService、AgentContextService、ToolRegistry、AgentRun 等参与类。")
    add_placeholder(doc, "图4-4 账单导入确认活动图", "展示上传截图、识别候选交易、用户确认、写入 transaction 和更新候选状态过程。")
    add_placeholder(doc, "图4-5 账单导入 VOPC 类图", "展示 BillImport.vue、BillImportController、BillImportServiceImpl、BillAiClient、BillImportRecord、BillCandidateTransaction、TransactionService 等参与类。")

    add_heading(doc, "4.3 实体类类图", 2)
    add_p(doc, "系统实体类围绕用户展开。User 与 FinancialProfile、Transaction、Budget、BudgetAlert、ChatConversation、AgentMemory、AgentSkill、AgentSchedule、PendingAction 等实体存在一对多或一对一关系。ChatConversation 与 ChatMessage 形成会话和消息关系；AgentRun 与 AgentRunStep 通过 traceId 关联；BillImportRecord 与 BillCandidateTransaction 表示一次账单识别及其候选交易。")
    add_placeholder(doc, "图4-6 实体类类图", "展示系统核心实体、主要属性和实体之间的一对一、一对多关系。")


def add_design(doc):
    add_heading(doc, "5 系统设计", 1)
    add_heading(doc, "5.1 功能模块", 2)
    add_p(doc, "系统按照业务职责划分为认证模块、财务数据模块、统计分析模块、账单导入模块、智能对话模块、Agent 能力模块、审计与周期任务模块、系统个性化模块。模块之间通过前后端接口和后端服务层协作，既能覆盖传统财务管理功能，也能支撑 Agent 化能力扩展。")
    add_placeholder(doc, "图5-1 功能模块图", "展示认证、交易、预算、统计、画像、账单导入、AI 对话、Skills、周期任务、审计和待确认动作等模块。")

    add_heading(doc, "5.2 技术选型与运行环境", 2)
    add_table(doc, ["层级", "技术方案", "说明"], [
        ("开发语言", "Java 17、JavaScript、Python", "Java 用于后端业务服务，JavaScript 用于 Vue 前端，Python 用于 AI 辅助服务和文档脚本。"),
        ("前端技术", "Vue 3、Vite、Pinia、Vue Router、Tailwind CSS、ECharts", "支持组件化开发、状态管理、页面路由、快速构建和统计图表展示。"),
        ("后端技术", "Spring Boot 3.2.5、MyBatis-Plus、JWT", "负责 REST 接口、业务逻辑、数据库访问和登录鉴权。"),
        ("数据库", "MySQL 8.0", "持久化用户、交易、预算、画像、聊天和 Agent 运行数据。"),
        ("AI 能力", "LangChain4j、DashScope、RAG、ReAct、Tavily Search", "提供智能问答、财务知识增强、工具调用和外部信息检索能力。"),
        ("开发工具", "IntelliJ IDEA、VS Code、Maven、npm", "支撑前后端编码、依赖管理和本地调试。"),
        ("建模工具", "PlantUML", "用于绘制用例图、活动图、类图、架构图和 ER 图。"),
    ], [2.6, 5.2, 7.2])
    add_table(doc, ["项目", "说明"], [
        ("操作系统", "Windows 开发环境，支持 Linux 服务器部署。"),
        ("Web 服务器", "Spring Boot 内嵌 Tomcat；前端开发阶段使用 Vite Dev Server。"),
        ("浏览器", "Chrome 或 Edge。"),
        ("JDK / Node 版本", "JDK 17+，Node.js 18+，npm 9+。"),
        ("数据库", "MySQL 8.0+，字符集 utf8mb4。"),
        ("默认端口", "后端 8080，前端 3000。"),
    ], [4, 11])

    add_heading(doc, "5.3 系统架构", 2)
    add_heading(doc, "5.3.1 架构设计（包图/分层图）", 3)
    add_placeholder(doc, "图5-2 系统架构分层图", "展示 Vue 前端、Spring Boot 后端、MyBatis-Plus 数据访问、MySQL 数据库、LangChain4j/DashScope/Tavily 外部能力。")
    add_heading(doc, "5.3.2 软件架构说明", 3)
    add_p(doc, "系统采用前后端分离架构。前端基于 Vue 3 构建单页应用，负责页面展示、表单交互、图表展示和 API 调用；后端基于 Spring Boot 提供 REST 与 SSE 接口，负责鉴权、业务处理、Agent 编排和数据持久化；数据库使用 MySQL 存储业务数据。")
    add_p(doc, "后端内部采用 Controller、Service、Mapper、Entity/DTO 的分层设计。Controller 层处理 HTTP 请求和用户身份上下文，Service 层实现业务流程，Mapper 层通过 MyBatis-Plus 访问数据库。Agent 能力作为独立业务层扩展，包含上下文组装、RAG 检索、ReAct 推理、工具注册、长期记忆、反思和周期任务。")

    add_heading(doc, "5.4 接口设计", 2)
    add_table(doc, ["模块", "接口", "方式", "说明"], [
        ("认证", "/api/auth/register", "POST", "用户注册，写入用户名和加密密码。"),
        ("认证", "/api/auth/login", "POST", "用户登录，返回 JWT 和用户信息。"),
        ("交易", "/api/transactions", "GET/POST", "查询或新增交易记录。"),
        ("交易", "/api/transactions/{id}", "GET/PUT/DELETE", "查看、修改或删除指定交易。"),
        ("统计", "/api/transactions/category-summary", "GET", "按分类汇总交易数据。"),
        ("预算", "/api/budgets", "GET/PUT", "查询或保存预算配置。"),
        ("财务画像", "/api/financial-profile", "GET/PUT", "读取或保存用户长期财务画像。"),
        ("账单导入", "/api/bills/import", "POST", "上传账单截图并生成识别记录。"),
        ("账单导入", "/api/bills/{id}/confirm", "POST", "确认候选交易并写入正式记录。"),
        ("聊天", "/api/chat/react/stream", "POST/SSE", "进入流式 ReAct 对话链路。"),
        ("Agent 运行", "/api/agent-runs/{traceId}", "GET", "根据 traceId 查询 Agent 运行步骤。"),
        ("Agent Skills", "/api/agent-skills", "GET", "查询 Agent 可用技能。"),
        ("周期任务", "/api/agent-schedules", "GET/POST", "查询或创建周期任务。"),
        ("待确认动作", "/api/pending-actions/{id}/confirm", "POST", "确认 AI 生成的写操作。"),
    ], [2.3, 4.7, 2.5, 5.5])

    add_heading(doc, "5.5 类设计", 2)
    add_p(doc, "类设计围绕核心业务职责展开。认证与用户模块包含 User、UserServiceImpl、AuthController；财务记录模块包含 Transaction、Budget、FinancialProfile 及其 Service；账单导入模块包含 BillImportRecord、BillCandidateTransaction、BillImportServiceImpl 和 BillAiClient；Agent 模块包含 ReActAgentService、AgentContextService、ToolRegistry、AgentRun、AgentRunStep、AgentMemory、AgentSkill、AgentSchedule 和 PendingAction。")
    add_placeholder(doc, "图5-3 详细设计类图", "展示主要 Controller、Service、Entity、Mapper 与 Agent 组件之间的依赖关系。")

    add_heading(doc, "5.6 数据库设计", 2)
    add_heading(doc, "5.6.1 数据模型（ER图）", 3)
    add_placeholder(doc, "图5-4 数据库 ER 图", "展示 user 与 transaction、budget、financial_profile、chat_conversation、agent_run、agent_skill、agent_schedule、bill_import_record 等表之间的关系。")
    add_heading(doc, "5.6.2 表结构设计", 3)
    add_table(doc, ["表名", "关键字段", "说明"], [
        ("user", "id、username、password、nickname、email、created_at、deleted", "用户基础信息表，username 唯一。"),
        ("financial_profile", "id、user_id、life_stage、monthly_income、risk_preference、savings_goal_amount", "用户长期财务画像，一名用户对应一份画像。"),
        ("transaction", "id、user_id、amount、type、category、description、transaction_date", "正式收入支出记录，是统计分析的基础。"),
        ("expense_category", "id、user_id、name、icon、benchmark_min、benchmark_max、sort_order", "消费分类表，支持系统默认分类和用户自定义分类。"),
        ("budget", "id、user_id、category、month、budget_amount、alert_threshold", "预算配置表，按用户、分类和月份唯一。"),
        ("budget_alert", "id、user_id、category、month、alert_type、severity、spent_amount", "预算预警记录表，用于提醒超支和风险趋势。"),
        ("chat_conversation", "id、user_id、title、created_at、updated_at", "聊天会话表，组织多轮对话。"),
        ("chat_message", "id、user_id、conversation_id、role、content、trace_id", "聊天消息表，记录用户和助手消息。"),
        ("agent_run", "id、user_id、trace_id、query、final_answer、status、duration_ms", "Agent 运行记录表，按 traceId 唯一追踪一次执行。"),
        ("agent_run_step", "id、user_id、trace_id、step_number、tool_name、input、status", "Agent 步骤表，记录每一步工具调用或执行状态。"),
        ("agent_memory", "id、user_id、memory_type、memory_key、memory_value、confidence、disabled", "Agent 长期记忆表，存储低风险偏好和长期指令。"),
        ("agent_skill", "id、user_id、skill_key、name、risk_level、instruction_text、enabled", "可安装 Agent Skill 表。"),
        ("pending_action", "id、user_id、action_type、title、payload、status", "待确认动作表，保证关键写操作先审后执。"),
        ("bill_import_record", "id、user_id、file_path、bill_type、confidence、status", "账单图片导入记录表。"),
        ("bill_candidate_transaction", "id、bill_import_id、user_id、amount、category、transaction_date、status", "账单识别出的候选交易表。"),
        ("agent_schedule", "id、user_id、name、cron_expression、task_query、enabled、next_run_at", "Agent 周期任务配置表。"),
        ("agent_schedule_run", "id、schedule_id、user_id、trace_id、status、answer、duration_ms", "周期任务运行历史表。"),
    ], [3.2, 6.5, 6.2])


def add_implementation(doc):
    add_heading(doc, "6 编码实现", 1)
    add_heading(doc, "6.1 项目结构", 2)
    add_code_block(doc, "项目目录结构", """
smart_finance_agent/
├── backend/                 # Spring Boot 后端
│   ├── src/main/java/       # Controller、Service、Entity、DTO、Mapper、Agent 逻辑
│   ├── src/main/resources/  # application.yml、schema.sql、data.sql、knowledge
│   └── src/test/            # 后端测试
├── frontend/                # Vue 3 前端
│   ├── src/api/             # API 请求封装
│   ├── src/components/      # 业务组件与通用组件
│   ├── src/layouts/         # 主布局
│   ├── src/router/          # 页面路由
│   └── src/views/           # 页面视图
├── env.example
├── start-dev.ps1
└── README.md
""")
    add_p(doc, "后端主要包路径为 com.smartfinance.agent，按 controller、service、service.impl、mapper、entity、dto、agent、context、config 等职责划分。前端主要页面位于 frontend/src/views，接口封装位于 frontend/src/api，路由配置位于 frontend/src/router/index.js。")

    add_heading(doc, "6.2 关键代码", 2)
    add_code_block(doc, "ChatController 流式对话入口", """
@PostMapping(value = "/react/stream", produces = "text/event-stream")
public SseEmitter reactStream(@RequestBody ChatRequest request) {
    Long userId = UserIdContext.getUserId();
    return chatService.chatStream(userId, request);
}
""")
    add_p(doc, "该接口是智能财务问答的主要入口。前端提交用户问题后，后端通过 SSE 返回流式结果，使用户能够看到 Agent 的运行过程和最终回答。")
    add_code_block(doc, "ReActAgentService 核心职责", """
AgentContext context = agentContextService.buildContext(userId, conversationId, query);
String ragContext = ragKnowledgeService.retrieve(query);
List<ToolSpec> tools = toolRegistry.getPromptVisibleTools(query);
AgentDecision decision = callModelWithJsonDecision(context, ragContext, tools);
ToolResult observation = toolRegistry.execute(decision.toolName(), decision.input());
agentRunService.recordStep(traceId, observation);
""")
    add_p(doc, "ReActAgentService 负责把用户问题、历史上下文、财务画像、长期记忆、RAG 知识和可见工具组织起来，并通过模型决策是否调用工具。每一步工具调用都会写入运行轨迹，便于审计。")
    add_code_block(doc, "ToolRegistry 工具执行边界", """
public ToolExecutionResult execute(Long userId, String toolName, String inputJson) {
    ToolDefinition tool = tools.get(toolName);
    if (tool == null) {
        return ToolExecutionResult.failed("未知工具");
    }
    return tool.executor().execute(userId, inputJson);
}
""")
    add_p(doc, "ToolRegistry 将模型可见的工具说明与实际可执行工具分离，避免模型直接越权调用未注册能力。记账、预算设置、创建任务等写操作会转化为待确认动作。")
    add_code_block(doc, "账单导入确认流程", """
BillImportRecord record = saveOriginalImage(userId, file);
BillAiResult result = billAiClient.analyze(file);
saveCandidateTransactions(record.getId(), userId, result);
// 用户确认后：
Transaction tx = convertCandidateToTransaction(candidate);
transactionService.create(userId, tx);
markCandidateConfirmed(candidate.getId(), tx.getId());
""")
    add_p(doc, "账单导入流程体现了系统的安全设计：AI 识别只生成候选数据，只有用户确认后才写入正式交易表。")

    add_heading(doc, "6.3 实现效果", 2)
    for idx, (name, path, desc) in enumerate(SCREENSHOTS, 1):
        if not path.exists():
            add_p(doc, f"图6-{idx} {name}：截图文件未找到，路径为 {path}")
            continue
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(path), width=Inches(6.25))
        add_caption(doc, f"图6-{idx} {name}")
        add_p(doc, desc)


def add_summary(doc):
    add_heading(doc, "7 总结与体会", 1)
    add_heading(doc, "7.1 实践收获", 2)
    add_p(doc, "通过本次实训，我完整经历了从系统分析、用例建模、架构设计、数据库设计到编码实现说明的过程。项目不只是实现若干页面和接口，而是要求把真实业务目标转化为参与者、用例、分析类、实体类、接口和表结构，这一过程加深了我对软件工程建模方法的理解。")
    add_p(doc, "在技术实践方面，本项目综合使用 Vue 3、Spring Boot、MyBatis-Plus、MySQL、JWT、LangChain4j、RAG 和 ReAct 等技术，进一步理解了前后端分离系统的分层方式，也学习了如何在传统业务系统中接入大模型能力，并通过审计、确认和持久化机制约束 AI 行为。")
    add_heading(doc, "7.2 遇到的问题及解决方法", 2)
    add_table(doc, ["问题", "原因分析", "解决方法"], [
        ("AI 生成内容可能直接改变用户数据", "大模型回答具有不确定性，若直接执行写操作会带来误记账、误设预算等风险。", "设计 pending_action 待确认动作机制，将关键写操作转为用户确认后执行。"),
        ("聊天上下文容易过长", "多轮对话、RAG 片段、工具结果和历史记录都会占用模型上下文窗口。", "通过 AgentContextService、ContextBudgetService 和上下文压缩摘要控制输入规模。"),
        ("账单截图识别结果不一定完全准确", "图片质量、账单来源和文字遮挡都会影响识别结果。", "将识别结果保存为候选交易，用户核对后再写入 transaction 表。"),
        ("Agent 执行过程难以解释", "如果只保存最终回答，无法定位工具调用和错误原因。", "引入 traceId、agent_run、agent_run_step、analysis_record 和 skill_invocation_record 记录运行轨迹。"),
    ], [4, 5.5, 5.5])
    add_heading(doc, "7.3 总结与展望", 2)
    add_p(doc, "目前系统已经实现登录注册、收支记录、预算预警、统计分析、财务画像、账单导入、AI 对话、长期记忆、Agent Skills、周期任务、审计和待确认动作等核心能力，能够支撑个人智能财务管理的主要流程。系统采用前后端分离和分层设计，业务模块边界较清晰，后续可继续扩展更多财务场景。")
    add_p(doc, "后续改进方向包括：完善股票分析模块和实时行情能力；增加更多自动化测试和端到端测试；优化移动端适配；增强账单识别的准确率和隐私保护；完善权限管理和部署方案；进一步细化 Agent 工具权限、超时、审计和安全策略。")


def build():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document(str(TEMPLATE)) if TEMPLATE.exists() else Document()
    setup_document(doc)
    if TEMPLATE.exists():
        trim_template_body_from(doc, "系统概述")
        fill_template_cover(doc)
    else:
        clear_document(doc)
        add_cover(doc)
        add_toc(doc)
    add_system_overview(doc)
    add_business_model(doc)
    add_use_cases(doc)
    add_analysis(doc)
    add_design(doc)
    add_implementation(doc)
    add_summary(doc)
    output_path = OUT_DOCX
    try:
        doc.save(str(output_path))
    except PermissionError:
        output_path = FALLBACK_DOCX
        doc.save(str(output_path))
    with zipfile.ZipFile(output_path) as zf:
        names = set(zf.namelist())
        assert "word/document.xml" in names
        media_count = len([n for n in names if n.startswith("word/media/")])
    print(f"wrote={output_path}")
    print(f"media_count={media_count}")
    print(f"size={output_path.stat().st_size}")


if __name__ == "__main__":
    build()
